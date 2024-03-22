import os
from docx import Document
from openpyxl import load_workbook
import PyPDF2
from datetime import datetime

def extract_docx_metadata(file_path):
    doc = Document(file_path)
    metadata = {}
    core_properties = doc.core_properties
    metadata['Title'] = core_properties.title
    metadata['Author'] = core_properties.author
    metadata['Subject'] = core_properties.subject
    metadata['Keywords'] = core_properties.keywords
    metadata['Comments'] = core_properties.comments
    return metadata

def extract_xlsx_metadata(file_path):
    wb = load_workbook(filename=file_path)
    metadata = {}
    props = wb.properties
    metadata['Title'] = props.title
    metadata['Author'] = props.creator
    metadata['Subject'] = props.subject
    metadata['Keywords'] = props.keywords
    metadata['Comments'] = props.description
    return metadata

def extract_pdf_metadata(file_path):
    pdf_file = open(file_path, 'rb')
    reader = PyPDF2.PdfReader(pdf_file)
    metadata = {}
    metadata['Title'] = reader.metadata.title
    metadata['Author'] = reader.metadata.author
    metadata['Subject'] = reader.metadata.subject
    
    if '/Keywords' in reader.metadata:
        metadata['Keywords'] = reader.metadata['/Keywords']
    else:
        metadata['Keywords'] = None

    return metadata


def extract_metadata(directory):
    metadata = {}
    for filename in os.listdir(directory):
        file_path = os.path.join(directory, filename)
        if filename.endswith('.docx'):
            metadata[filename] = extract_docx_metadata(file_path)
        elif filename.endswith('.xlsx'):
            metadata[filename] = extract_xlsx_metadata(file_path)
        elif filename.endswith('.pdf'):
            metadata[filename] = extract_pdf_metadata(file_path)
    return metadata

def main():
    directory = os.path.join("C:\\Users\\diego\\Desktop\\Cyber\\Files")
    metadata = extract_metadata(directory)

    for filename, data in metadata.items():
        print("Metadata for:", filename)
        for key, value in data.items():
            print(key + ':', value)
        print()

if __name__ == "__main__":
    main()
